import hashlib
import hmac
import io
import re
import random
import secrets
from datetime import timedelta

from django.conf import settings
from django.core.mail import EmailMessage, send_mail
from django.db import transaction
from django.utils import timezone

from .models import (LEVELS, SKILLS, AssessmentSetting, Attempt, AttemptAnswer, AuditLog,
                     EmailOTP, Question, ReportLog)

_rng = random.SystemRandom()


class NotEnoughQuestions(Exception):
    pass


def audit(user, action, details=""):
    AuditLog.objects.create(actor=user if getattr(user, "is_authenticated", False) else None,
                            action=action, details=details[:2000])


# ---------------------------------------------------------------- attempts
@transaction.atomic
def start_attempt(student):
    """Return existing in-progress attempt or create a new random one. None if limit reached."""
    cfg = AssessmentSetting.get()
    running = student.attempts.filter(status="in_progress").first()
    if running:
        return running
    if student.attempts.count() >= cfg.max_attempts:
        return None

    n = cfg.questions_per_section
    picked = []
    for skill, _ in SKILLS:
        for level, _ in LEVELS:
            ids = list(Question.objects.filter(skill=skill, difficulty=level, active=True)
                       .values_list("id", flat=True))
            if len(ids) < n:
                raise NotEnoughQuestions(f"Not enough active {skill}/{level} questions ({len(ids)}/{n}).")
            picked.extend(_rng.sample(ids, n))  # sample = no duplicates

    now = timezone.now()
    attempt = Attempt.objects.create(
        student=student, started_at=now,
        deadline=now + timedelta(minutes=cfg.time_limit_minutes) if cfg.time_limit_minutes else None,
    )
    rows = []
    for pos, qid in enumerate(picked, start=1):
        letters = list("ABCD")
        _rng.shuffle(letters)
        rows.append(AttemptAnswer(attempt=attempt, question_id=qid, position=pos, option_order="".join(letters)))
    AttemptAnswer.objects.bulk_create(rows)
    audit(student.user, "attempt_started", f"attempt {attempt.pk}")
    return attempt


def exam_sections(attempt):
    """Structure for the exam page. NEVER includes answer keys."""
    answers = list(attempt.answers.select_related("question"))
    sections = []
    for skill, skill_label in SKILLS:
        levels = []
        for level, level_label in LEVELS:
            items = []
            for a in answers:
                q = a.question
                if q.skill == skill and q.difficulty == level:
                    opts = [{"label": "ABCD"[i], "value": orig, "text": q.option_text(orig)}
                            for i, orig in enumerate(a.option_order)]
                    items.append({"answer": a, "question": q, "options": opts})
            levels.append({"key": level, "label": level_label, "items": items})
        sections.append({"key": skill, "label": skill_label, "levels": levels,
                         "count": sum(len(l["items"]) for l in levels)})
    return sections


def evaluate(setting, sql_s, sql_m, py_s, py_m):
    total_s, total_m = sql_s + py_s, sql_m + py_m
    pct = round(total_s * 100 / total_m, 2) if total_m else 0
    need = setting.pass_percentage
    if setting.pass_mode == "each_skill":
        ok = (sql_m and sql_s * 100 >= need * sql_m) and (py_m and py_s * 100 >= need * py_m)
    else:
        ok = total_m and total_s * 100 >= need * total_m
    return pct, "PASS" if ok else "FAIL"


@transaction.atomic
def grade_attempt(attempt):
    attempt = Attempt.objects.select_for_update().get(pk=attempt.pk)
    if attempt.status == "submitted":
        return attempt
    breakdown = {s: {l: {"score": 0, "max": 0} for l, _ in LEVELS} for s, _ in SKILLS}
    changed = []
    for a in attempt.answers.select_related("question"):
        q = a.question
        ok = bool(a.selected) and a.selected == q.answer_key
        a.is_correct = ok
        a.marks_awarded = q.marks if ok else 0
        changed.append(a)
        cell = breakdown[q.skill][q.difficulty]
        cell["max"] += q.marks
        cell["score"] += a.marks_awarded
    AttemptAnswer.objects.bulk_update(changed, ["is_correct", "marks_awarded"])

    tot = lambda s, k: sum(v[k] for v in breakdown[s].values())
    attempt.sql_score, attempt.sql_max = tot("sql", "score"), tot("sql", "max")
    attempt.python_score, attempt.python_max = tot("python", "score"), tot("python", "max")
    attempt.total_score = attempt.sql_score + attempt.python_score
    attempt.total_max = attempt.sql_max + attempt.python_max
    attempt.percentage, attempt.result = evaluate(
        AssessmentSetting.get(), attempt.sql_score, attempt.sql_max, attempt.python_score, attempt.python_max)
    attempt.breakdown = breakdown
    attempt.status = "submitted"
    attempt.submitted_at = timezone.now()
    attempt.save()
    audit(attempt.student.user, "attempt_submitted",
          f"attempt {attempt.pk}: {attempt.total_score}/{attempt.total_max} {attempt.result}")
    return attempt


def result_summary(attempt):
    out = []
    for skill, label in SKILLS:
        b = attempt.breakdown.get(skill, {})
        score = getattr(attempt, f"{skill}_score")
        mx = getattr(attempt, f"{skill}_max")
        out.append({
            "label": label, "score": score, "max": mx,
            "pct": round(score * 100 / mx, 1) if mx else 0,
            "levels": [{"label": ll, "score": b.get(l, {}).get("score", 0), "max": b.get(l, {}).get("max", 0)}
                       for l, ll in LEVELS],
        })
    return out


# ---------------------------------------------------------------- email report
def report_recipients():
    """Configured report address + every active admin (host) user's email, no duplicates."""
    from django.contrib.auth.models import User
    emails = [AssessmentSetting.get().report_recipient]
    emails += list(User.objects.filter(is_staff=True, is_active=True).exclude(email="")
                   .values_list("email", flat=True))
    seen, out = set(), []
    for e in emails:
        if e and e.lower() not in seen:
            seen.add(e.lower())
            out.append(e)
    return out


XLSX_MIME = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"


def build_report_xlsx(attempt):
    """Build the candidate's exam report as an Excel workbook. Returns bytes.

    Sheet 1 "Report"        - one table row: columns for candidate details, scores, result.
    Sheet 2 "Answer Detail" - every question with the candidate's answer, the correct answer and marks.
    """
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

    s = attempt.student
    font = Font(name="Arial", size=10)
    bold = Font(name="Arial", size=10, bold=True)
    head_font = Font(name="Arial", size=10, bold=True, color="FFFFFF")
    head_fill = PatternFill("solid", fgColor="1F3864")
    label_fill = PatternFill("solid", fgColor="DDEBF7")
    green = PatternFill("solid", fgColor="C6EFCE")
    red = PatternFill("solid", fgColor="FFC7CE")
    thin = Side(style="thin", color="BFBFBF")
    box = Border(left=thin, right=thin, top=thin, bottom=thin)
    wrap = Alignment(wrap_text=True, vertical="top")

    wb = Workbook()

    # ------------------------------------------------------------ Report (one table row, like a DB table)
    ws = wb.active
    ws.title = "Report"
    submitted = (timezone.localtime(attempt.submitted_at).strftime("%d-%m-%Y %H:%M")
                 if attempt.submitted_at else "")
    columns = [
        ("S.No", 1, 7),
        ("Attempt ID", attempt.pk, 12),
        ("Candidate Name", s.name, 24),
        ("Roll No", s.roll_number or "", 14),
        ("Email", s.email, 30),
        ("Phone", s.phone or "", 15),
        ("Submitted At", submitted, 18),
        ("SQL Score", attempt.sql_score, 11),
        ("SQL Max", attempt.sql_max, 10),
        ("Python Score", attempt.python_score, 13),
        ("Python Max", attempt.python_max, 12),
        ("Total Score", attempt.total_score, 12),
        ("Total Max", attempt.total_max, 11),
        ("Percentage", float(attempt.percentage), 12),
        ("Result", attempt.result, 10),
        ("Tab Switches", attempt.tab_switches, 13),
        ("Submission Type", "Auto-submitted" if attempt.auto_submitted else "Normal", 17),
    ]
    for c, (head, value, width) in enumerate(columns, start=1):
        h = ws.cell(row=1, column=c, value=head)
        h.font, h.fill, h.border = head_font, head_fill, box
        h.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        v = ws.cell(row=2, column=c, value=value)
        v.font, v.border = font, box
        v.alignment = Alignment(horizontal="center" if isinstance(value, (int, float)) else "left",
                                vertical="center")
        ws.column_dimensions[ws.cell(row=1, column=c).column_letter].width = width
    ws.cell(row=2, column=14).number_format = "0.00"
    res = ws.cell(row=2, column=15)
    res.fill = green if attempt.result == "PASS" else red
    res.font = bold
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:{ws.cell(row=1, column=len(columns)).column_letter}2"

    # ------------------------------------------------------------ Answer Detail
    wd = wb.create_sheet("Answer Detail")
    headers = ["S.No", "Skill", "Level", "Question", "Candidate Answer", "Correct Answer", "Result", "Marks"]
    for c, h in enumerate(headers, start=1):
        cell = wd.cell(row=1, column=c, value=h)
        cell.font, cell.fill, cell.border = head_font, head_fill, box
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

    def opt(q, letter):
        return f"{letter}. {q.option_text(letter)}" if letter else "Not answered"

    r = 2
    for a in attempt.answers.select_related("question"):
        q = a.question
        if not a.selected:
            status = "Not answered"
        else:
            status = "Correct" if a.is_correct else "Wrong"
        values = [a.position, q.get_skill_display(), q.get_difficulty_display(), q.prompt,
                  opt(q, a.selected), opt(q, q.answer_key), status, a.marks_awarded]
        for c, v in enumerate(values, start=1):
            cell = wd.cell(row=r, column=c, value=v)
            cell.font, cell.border, cell.alignment = font, box, wrap
        fill = green if status == "Correct" else red
        wd.cell(row=r, column=7).fill = fill
        r += 1
    for col, width in zip("ABCDEFGH", (6, 10, 10, 60, 38, 38, 14, 8)):
        wd.column_dimensions[col].width = width
    wd.freeze_panes = "A2"
    wd.auto_filter.ref = f"A1:H{max(r - 1, 1)}"

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def send_report(attempt, log=None):
    """Email the exam report to the recipients as an Excel (.xlsx) attachment."""
    recipients = report_recipients()
    if log is None:
        log = ReportLog.objects.create(attempt=attempt, recipient=", ".join(recipients)[:254])
    s = attempt.student
    body = (
        f"Please find the assessment report of {s.name} ({s.roll_number}) attached as an Excel file.\n\n"
        f"Result: {attempt.result} - {attempt.total_score}/{attempt.total_max} ({attempt.percentage}%)\n"
    )
    try:
        safe = re.sub(r"[^A-Za-z0-9_-]+", "_", f"{s.roll_number or s.pk}_{s.name}").strip("_")
        msg = EmailMessage(f"Assessment report - {s.name} - {attempt.result}", body,
                           settings.DEFAULT_FROM_EMAIL, recipients)
        msg.attach(f"Assessment_Report_{safe}_{attempt.pk}.xlsx", build_report_xlsx(attempt), XLSX_MIME)
        msg.send(fail_silently=False)
        log.delivery_status, log.sent_at, log.error = "sent", timezone.now(), ""
    except Exception as exc:  # logged, retried later
        log.delivery_status, log.error = "failed", str(exc)[:1000]
        log.retries += 1
    log.save()
    return log


# ---------------------------------------------------------------- OTP
def _hash(user, code):
    return hashlib.sha256(f"{settings.SECRET_KEY}:{user.pk}:{code}".encode()).hexdigest()


def issue_otp(user):
    EmailOTP.objects.filter(user=user, used=False).update(used=True)
    code = f"{secrets.randbelow(10 ** 6):06d}"
    EmailOTP.objects.create(user=user, code_hash=_hash(user, code))
    send_mail("Your ITA Associates verification code",
              f"Your one-time code is {code}. It is valid for {settings.OTP_VALID_MINUTES} minutes.\n"
              "If you did not try to sign in, ignore this email.",
              settings.DEFAULT_FROM_EMAIL, [user.email], fail_silently=False)


def check_otp(user, code):
    otp = EmailOTP.objects.filter(user=user, used=False).order_by("-created_at").first()
    if not otp:
        return False
    if timezone.now() - otp.created_at > timedelta(minutes=settings.OTP_VALID_MINUTES) \
            or otp.tries >= settings.OTP_MAX_TRIES:
        otp.used = True
        otp.save(update_fields=["used"])
        return False
    otp.tries += 1
    if hmac.compare_digest(otp.code_hash, _hash(user, code.strip())):
        otp.used = True
        otp.save(update_fields=["used", "tries"])
        return True
    otp.save(update_fields=["tries"])
    return False